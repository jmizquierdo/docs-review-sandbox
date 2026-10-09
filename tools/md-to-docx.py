#!/usr/bin/env python3
"""Render a Markdown document to a formatted Word file.

Pandoc does the conversion; the formatting rules in a YAML file
(default: docx-format.yaml at the repository root) are then applied
to the Word file. The Markdown source is never modified.

Usage:
    python tools/md-to-docx.py <input.md> [output.docx] [--config docx-format.yaml]
"""

import argparse
import re
import sys
import tempfile
from pathlib import Path

import pypandoc
import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_CONFIG = REPO_ROOT / "docx-format.yaml"

HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
FIGURE_CAPTION_RE = re.compile(r"^\*Figure \d+\.[^*]*\*\s*$")
TABLE_CAPTION_RE = re.compile(r"^\*Table \d+\.[^*]*\*\s*$")


def slugify(text: str) -> str:
    """GitHub-style heading slug: lowercase, drop punctuation, spaces to hyphens."""
    slug = re.sub(r"[^\w\s-]", "", text.lower()).strip()
    return re.sub(r"\s", "-", slug)


def prepare_markdown(content: str, cfg: dict) -> str:
    """Make in-document links work in Word and drop duplicate caption lines."""
    link_targets = set(re.findall(r"\]\(#([^)\s]+)\)", content))
    # Pandoc drops raw HTML in DOCX, so turn <a id="x"></a> into a bookmark span.
    content = re.sub(r'<a id="([^"]+)"></a>', r"[]{#\1}", content)

    strip_figures = cfg.get("figure_captions", {}).get("strip_italic_lines", True)
    strip_tables = cfg.get("table_captions", {}).get("strip_italic_lines", False)

    out = []
    in_code = False
    for line in content.split("\n"):
        if line.lstrip().startswith(("```", "~~~")):
            in_code = not in_code
        if not in_code:
            if strip_figures and FIGURE_CAPTION_RE.match(line):
                continue
            if strip_tables and TABLE_CAPTION_RE.match(line):
                continue
            m = HEADING_RE.match(line)
            # Pandoc's own heading IDs drop leading numbers; pin the GitHub-style ID
            # when a link in the document points to it.
            if m and "{#" not in line and slugify(m.group(2)) in link_targets:
                line = f"{m.group(1)} {m.group(2)} {{#{slugify(m.group(2))}}}"
        out.append(line)
    return "\n".join(out)


def add_page_break_before(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    el = OxmlElement("w:pageBreakBefore")
    el.set(qn("w:val"), "true")
    ppr.append(el)


def shade(paragraph, color: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), color)
    ppr.append(el)


def set_font_size(paragraph, size_pt) -> None:
    for run in paragraph.runs:
        run.font.size = Pt(size_pt)


def is_table_caption(paragraph) -> bool:
    runs = [r for r in paragraph.runs if r.text.strip()]
    return (
        bool(re.match(r"^Table \d+\.", paragraph.text.strip()))
        and bool(runs)
        and all(r.italic for r in runs)
    )


def apply_formatting(docx_path: Path, cfg: dict) -> None:
    breaks = cfg.get("page_breaks", {})
    h1_break = breaks.get("h1_page_break", True)
    h1_no_break = set(breaks.get("h1_no_break") or [])
    h2_break = set(breaks.get("h2_page_break") or [])

    center_images = cfg.get("images", {}).get("center", False)
    fig = cfg.get("figure_captions", {})
    tab = cfg.get("table_captions", {})
    code = cfg.get("code_blocks", {})

    doc = Document(str(docx_path))
    for p in doc.paragraphs:
        style = p.style.name if p.style is not None else ""
        text = p.text.strip()

        if (style == "Heading 1" and h1_break and text not in h1_no_break) or (
            style == "Heading 2" and text in h2_break
        ):
            add_page_break_before(p)

        if center_images and p._p.findall(f"{qn('w:r')}/{qn('w:drawing')}"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if style == "Image Caption":
            if fig.get("center"):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if fig.get("font_size_pt"):
                set_font_size(p, fig["font_size_pt"])

        if is_table_caption(p):
            if tab.get("center"):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if tab.get("font_size_pt"):
                set_font_size(p, tab["font_size_pt"])

        if style == "Source Code":
            if code.get("background_color"):
                shade(p, code["background_color"])
            if code.get("font_size_pt"):
                set_font_size(p, code["font_size_pt"])

    if cfg.get("tables", {}).get("center"):
        for table in doc.tables:
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

    doc.save(str(docx_path))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("input", type=Path, help="Markdown file to render")
    parser.add_argument("output", type=Path, nargs="?", help="Word file to write (default: next to the input)")
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG, help="YAML formatting rules")
    args = parser.parse_args()

    if not args.input.is_file():
        print(f"Error: {args.input} not found", file=sys.stderr)
        return 1
    output = args.output or args.input.with_suffix(".docx")

    cfg = {}
    if args.config.is_file():
        cfg = yaml.safe_load(args.config.read_text(encoding="utf-8")) or {}
        print(f"Using config: {args.config}")
    else:
        print(f"No config at {args.config}, using defaults")

    content = prepare_markdown(args.input.read_text(encoding="utf-8"), cfg)

    # The temporary copy sits next to the source so relative image paths still resolve.
    with tempfile.NamedTemporaryFile(
        "w", suffix=".md", dir=args.input.parent, delete=False, encoding="utf-8"
    ) as tmp:
        tmp.write(content)
    tmp_path = Path(tmp.name)
    try:
        output.parent.mkdir(parents=True, exist_ok=True)
        pypandoc.convert_file(
            str(tmp_path),
            "docx",
            format="markdown",
            outputfile=str(output),
            extra_args=[f"--resource-path={args.input.parent}"],
        )
    finally:
        tmp_path.unlink(missing_ok=True)

    apply_formatting(output, cfg)
    print(f"Created: {output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
